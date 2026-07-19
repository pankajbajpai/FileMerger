"""Merge multiple .docx files into a single .docx, with optional per-file passwords."""

import io

import msoffcrypto
from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer


class DocxMergeError(Exception):
    pass


def _is_encrypted(stream):
    stream.seek(0)
    try:
        office_file = msoffcrypto.OfficeFile(stream)
        return office_file.is_encrypted()
    except Exception:
        return False
    finally:
        stream.seek(0)


def _decrypt(stream, password, name):
    stream.seek(0)
    office_file = msoffcrypto.OfficeFile(stream)
    try:
        office_file.load_key(password=password)
        decrypted = io.BytesIO()
        office_file.decrypt(decrypted)
        decrypted.seek(0)
        return decrypted
    except Exception as exc:
        raise DocxMergeError(f"Incorrect password for '{name}'.") from exc


def merge_docx(files):
    """
    files: list of dicts:
        {
            "stream": file-like object,
            "name": original filename,
            "password": str or None,
        }

    Returns: BytesIO containing the merged .docx file.
    """
    if not files:
        raise DocxMergeError("No files supplied.")

    prepared = []
    for f in files:
        stream = f["stream"]
        name = f.get("name", "file.docx")
        password = f.get("password")

        stream.seek(0)
        if _is_encrypted(stream):
            if not password:
                raise DocxMergeError(
                    f"'{name}' is password protected. Please supply its password."
                )
            stream = _decrypt(stream, password, name)

        try:
            doc = Document(stream)
        except Exception as exc:
            raise DocxMergeError(f"Could not read '{name}': {exc}") from exc
        prepared.append((name, doc))

    master_name, master_doc = prepared[0]
    composer = Composer(master_doc)

    for name, doc in prepared[1:]:
        # Force each appended document to start on a new page.
        if doc.paragraphs:
            first_para = doc.paragraphs[0]
            run = first_para.insert_paragraph_before().add_run()
            run.add_break(WD_BREAK.PAGE)
        try:
            composer.append(doc)
        except Exception as exc:
            raise DocxMergeError(f"Could not merge '{name}': {exc}") from exc

    out = io.BytesIO()
    composer.save(out)
    out.seek(0)
    return out
